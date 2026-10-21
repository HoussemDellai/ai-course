import asyncio
import json
import shutil
import subprocess
import time
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from pydantic import ValidationError

from fakes import FakeNarrator, requires_ffmpeg
from generate import parse_request
from video_platform import documents, media
from video_platform.api import create_app
from video_platform.config import Settings
from video_platform.document_schemas import PageNarration
from video_platform.documents import (
    DocumentError, FoundryDocumentTeam, build_document_workflow, pdf_pages, powerpoint_pages, validate_document,
)
from video_platform.jobs import JobManager
from video_platform.schemas import JobStatus, ShotPrompt, VideoRequest
from video_platform.storage import LocalArtifactStore, read_json
from video_platform.video_models import RESOLUTIONS
from video_platform.workflow import PipelineDeps


def make_pdf(texts=("First page", "Second page")):
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Count {len(texts)} /Kids [{' '.join(f'{4 + i * 2} 0 R' for i in range(len(texts)))}] >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for i, text in enumerate(texts):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 320 180] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + i * 2} 0 R >>".encode()
        )
        color = "1 0 0" if i % 2 == 0 else "0 0 1"
        content = f"{color} rg 0 0 320 180 re f 0 0 0 rg BT /F1 14 Tf 20 90 Td ({text}) Tj ET".encode()
        objects.append(f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"\nendstream")
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    data.extend(b"".join(f"{offset:010} 00000 n \n".encode() for offset in offsets[1:]))
    data.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return bytes(data)


def make_pptx():
    deck = Presentation()
    for i in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = f"Slide {i + 1}"
        slide.notes_slide.notes_text_frame.text = f"Speaker notes for slide {i + 1}"
    out = BytesIO()
    deck.save(out)
    return out.getvalue()


class DocumentTeam:
    def __init__(self):
        self.pages = []

    async def narrate(self, page, prompt, language, seconds):
        self.pages.append((page, prompt, language, seconds))
        return PageNarration(title=f"Page {page.number}", narration=f"An explanation of page {page.number}.")


class Narrator(FakeNarrator):
    async def synthesize(self, *args, transport="sdk", **kwargs):
        assert transport == "rest"
        return await super().synthesize(*args, transport=transport, **kwargs)


@pytest.fixture
def setup(tmp_path, monkeypatch):
    monkeypatch.setitem(RESOLUTIONS, "1080p", (160, 90))
    settings = Settings(speech_endpoint="https://speech.example", work_dir=tmp_path / "work",
                        local_output_dir=tmp_path / "out", output_fps=12, max_document_mb=1,
                        libreoffice_path="missing-office-for-test")
    store = LocalArtifactStore(settings.local_output_dir)
    team, narrator = DocumentTeam(), Narrator(2.25)

    def factory(manager):
        deps = PipelineDeps(settings=settings, team=None, comfy=None, narrator=narrator,
                            store=store, progress=manager.progress, ops=manager.recorder)
        return build_document_workflow(deps, team)

    async def close():
        pass

    return settings, store, team, narrator, factory, close


def request(**kwargs):
    return VideoRequest(prompt="Explain this document to a beginner", use_case="document",
                        duration_minutes=0.25, **kwargs)


def test_document_schema_defaults_and_mode_guards():
    req = request()
    assert req.resolution == "1080p" and req.document.language == "en-US"
    assert VideoRequest(prompt="A movie").model_dump().get("document") is None
    assert "document" not in VideoRequest(prompt="A movie").model_dump()
    for values in (
        {"narration": False}, {"naturalistic": True}, {"music": True}, {"reference_image": True},
        {"upscaler": "seedvr2"}, {"education": {"public_topic": "Public data"}},
        {"document": {"language": "en"}}, {"document": {"language": "  "}},
    ):
        with pytest.raises(ValidationError):
            request(**values)
    with pytest.raises(ValidationError, match="Document options"):
        VideoRequest(prompt="A movie", document={})
    assert request(delivery={"rate_percent": 5}).delivery.rate_percent == 5


def test_pdf_validation_extracts_and_renders_in_order(tmp_path):
    data = make_pdf()
    upload = validate_document(data, r"C:\fakepath\Slides.PDF", len(data))
    assert upload.filename == "Slides.PDF" and upload.key == "input/document.pdf"
    pages = pdf_pages(data, tmp_path, (320, 180))
    assert [p.text for p in pages] == ["First page", "Second page"]
    assert [p.number for p in pages] == [1, 2]
    with Image.open(tmp_path / "page_000.png") as first, Image.open(tmp_path / "page_001.png") as second:
        assert first.size == second.size == (320, 180)
        assert first.getpixel((2, 2)) == (255, 0, 0)
        assert second.getpixel((2, 2)) == (0, 0, 255)


@pytest.mark.parametrize("name,data,code", [
    ("slides.ppt", b"not supported", 415), ("slides.docx", b"PK", 415),
    ("slides.pdf", b"not PDF", 422), ("slides.pdf", b"", 422), ("slides.pptx", b"PK broken", 422),
])
def test_invalid_uploads(name, data, code):
    with pytest.raises(DocumentError) as e:
        validate_document(data, name, 100000)
    assert e.value.status_code == code


def test_document_limits():
    with pytest.raises(DocumentError, match="larger") as e:
        validate_document(make_pdf(), "big.pdf", 5)
    assert e.value.status_code == 413
    with pytest.raises(DocumentError, match="1-60"):
        validate_document(make_pdf(["Page"] * 61), "many.pdf", 1000000)


def test_powerpoint_notes_and_unsafe_packages():
    data = make_pptx()
    assert validate_document(data, "deck.pptx", len(data)).key == "input/document.pptx"
    assert [p.notes for p in powerpoint_pages(data)] == [
        "Speaker notes for slide 1", "Speaker notes for slide 2",
    ]
    out = BytesIO(data)
    with ZipFile(out, "a") as archive:
        archive.writestr("ppt/slides/_rels/unsafe.xml.rels",
                         '<Relationships><Relationship TargetMode="External" Type="image" '
                         'Target="https://example.invalid/private.png"/></Relationships>')
    with pytest.raises(DocumentError, match="linked external"):
        powerpoint_pages(out.getvalue())


def test_powerpoint_with_embedded_chart_data_is_supported():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    data = CategoryChartData()
    data.categories = ["Before", "After"]
    data.add_series("Growth", (2, 4))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(5), Inches(4), data)
    out = BytesIO()
    deck.save(out)
    assert len(powerpoint_pages(out.getvalue())) == 1


def test_document_cli():
    req, source = parse_request([
        "Explain slides", "--use-case", "document", "--document", "deck.pptx",
        "--language", "fr-FR", "--voice", "fr-FR-DeniseNeural", "--speech-rate", "5",
    ])
    assert source == Path("deck.pptx") and req.document.language == "fr-FR"
    assert req.resolution == "1080p" and req.voice == "fr-FR-DeniseNeural"
    assert req.delivery.rate_percent == 5
    cutaways, _ = parse_request(["Explain slides", "--use-case", "document", "--document", "deck.pdf",
                                "--cinematic-cutaways", "--model", "ltx25"])
    assert cutaways.document.cinematic_cutaways and cutaways.video_model == "ltx25"
    assert not cutaways.document.animations
    animated, _ = parse_request(["Explain slides", "--use-case", "document", "--document", "deck.pdf",
                                 "--animations"])
    assert animated.document.animations and not animated.document.cinematic_cutaways
    for args in (
        ["--use-case", "document"], ["--document", "deck.pdf"],
        ["--language", "fr-FR"], ["--cinematic-cutaways"], ["--animations"],
        ["--use-case", "document", "--document", "deck.pdf", "--image", "a.png"],
        ["--use-case", "document", "--document", "deck.pdf", "--no-narration"],
        ["--use-case", "document", "--document", "deck.pdf", "--public-topic", "Public data"],
    ):
        with pytest.raises(SystemExit):
            parse_request(["Explain slides", *args])


@requires_ffmpeg
@pytest.mark.parametrize("kind,orientation", [("pdf", "horizontal"), ("pdf", "vertical"), ("pptx", "horizontal")])
async def test_document_pipeline_preserves_pages_voice_timing_and_retries(setup, monkeypatch, kind, orientation):
    settings, store, team, narrator, factory, close = setup
    if kind == "pptx":
        async def convert(source, settings):
            dest = source.with_suffix(".pdf")
            dest.write_bytes(make_pdf())
            return dest
        monkeypatch.setattr(documents, "convert_powerpoint", convert)
        monkeypatch.setattr("video_platform.jobs.office_available", lambda settings: True)
    data = make_pptx() if kind == "pptx" else make_pdf()
    upload = validate_document(data, f"deck.{kind}", settings.max_document_bytes)
    manager = JobManager(settings, store, factory)
    state = await manager.create(request(orientation=orientation), document=upload)
    await manager.wait(state.id)
    finished = await manager.get(state.id)
    assert finished.status == "completed", finished.error
    assert finished.clips_done == finished.clips_total == 2
    assert finished.request.document.filename == f"deck.{kind}"
    assert await store.exists(state.id, upload.key)
    assert [p[0].text for p in team.pages] == ["First page", "Second page"]
    if kind == "pptx":
        assert [p[0].notes for p in team.pages] == ["Speaker notes for slide 1", "Speaker notes for slide 2"]
    final = store.path(state.id, "final.mp4")
    assert await media.video_size(final) == ((160, 90) if orientation == "horizontal" else (90, 160))
    seconds, has_audio = await media.probe(final)
    assert has_audio and abs(seconds - 15) < 0.1
    levels = subprocess.run(
        ["ffmpeg", "-i", str(final), "-af", "volumedetect", "-f", "null", "-"],
        capture_output=True, text=True, check=True,
    )
    assert "max_volume:" in levels.stderr and "max_volume: -inf" not in levels.stderr
    assert narrator.texts == ["An explanation of page 1.", "An explanation of page 2."]
    timing = await read_json(store, state.id, "document/timing.json")
    assert timing["durations"] == [7.5, 7.5]
    for i, stamp in enumerate([1, 9]):
        frame = final.parent / f"check_{i}.png"
        await media.run("ffmpeg", "-y", "-ss", str(stamp), "-i", str(final), "-frames:v", "1", str(frame))
        with Image.open(frame) as image:
            r, g, b = image.getpixel((image.width // 2, image.height // 2))
            assert (r > 220 and b < 30) if i == 0 else (b > 220 and r < 30)
        frame.unlink()
    # Simulate an interrupted run and a replica without any local workspace.
    assert not (settings.work_dir / state.id).exists()
    await manager.progress(state.id, JobStatus.failed, error="Interrupted")
    await manager.retry(state.id)
    await manager.wait(state.id)
    assert (await manager.get(state.id)).status == "completed"
    assert len(team.pages) == len(narrator.texts) == 2
    assert any(op.status == "reused" for op in await manager.operations(state.id))
    await manager.shutdown()


@requires_ffmpeg
async def test_overlong_narration_fails_before_render(setup, monkeypatch):
    settings, store, team, narrator, factory, close = setup
    original_probe = media.probe

    async def too_long(path):
        if path.suffix == ".wav":
            return 301, True
        return await original_probe(path)

    monkeypatch.setattr(media, "probe", too_long)
    manager = JobManager(settings, store, factory)
    state = await manager.create(request(), document=validate_document(make_pdf(), "deck.pdf", 10000))
    await manager.wait(state.id)
    finished = await manager.get(state.id)
    assert finished.status == "failed" and "exceeding 10 minutes" in finished.error
    assert not await store.exists(state.id, "scenes/scene_000.mp4")
    await manager.shutdown()


@requires_ffmpeg
async def test_partial_failure_restores_pages_and_rejects_changed_source(setup, monkeypatch):
    settings, store, team, narrator, factory, close = setup
    mix = media.mix_narration

    async def fail_second(video, narration, dest, **kwargs):
        if dest.name == "scene_001.mp4":
            raise media.MediaError("Interrupted page render")
        return await mix(video, narration, dest, **kwargs)

    monkeypatch.setattr(media, "mix_narration", fail_second)
    manager = JobManager(settings, store, factory)
    state = await manager.create(request(), document=validate_document(make_pdf(), "deck.pdf", 10000))
    await manager.wait(state.id)
    assert (await manager.get(state.id)).status == "failed"
    assert await store.exists(state.id, "scenes/scene_000.mp4")
    assert not await store.exists(state.id, "scenes/scene_001.mp4")
    shutil.rmtree(settings.work_dir / state.id)
    monkeypatch.setattr(media, "mix_narration", mix)
    await manager.retry(state.id)
    await manager.wait(state.id)
    assert (await manager.get(state.id)).status == "completed"
    assert len(team.pages) == len(narrator.texts) == 2
    store.path(state.id, "input/document.pdf").write_bytes(make_pdf(["Different", "Source"]))
    await manager.progress(state.id, JobStatus.failed, error="Interrupted")
    await manager.retry(state.id)
    await manager.wait(state.id)
    assert "dependencies changed" in (await manager.get(state.id)).error
    assert len(team.pages) == len(narrator.texts) == 2
    await manager.shutdown()


async def test_image_only_page_and_missing_document_fail_explicitly(setup):
    settings, store, team, narrator, factory, close = setup
    manager = JobManager(settings, store, factory)
    with pytest.raises(DocumentError, match="one document"):
        await manager.create(request())
    state = await manager.create(request(), document=validate_document(make_pdf([""]), "scan.pdf", 10000))
    await manager.wait(state.id)
    finished = await manager.get(state.id)
    assert finished.status == "failed" and "no extractable text" in finished.error
    assert not team.pages and not narrator.texts
    await manager.shutdown()


def test_document_api_validation_and_private_artifacts(setup):
    settings, store, team, narrator, factory, close = setup
    settings = replace(settings, api_key="test-key")
    app = create_app(settings, services=(store, factory, close))
    headers = {"X-API-Key": "test-key"}
    payload = {"request": request(document={"filename": "spoofed.pptx"}).model_dump_json()}
    with TestClient(app) as client:
        def post(files=None, data=payload, **kwargs):
            return client.post("/api/videos/document", headers=headers, files=files, data=data, **kwargs)

        assert client.post("/api/videos/document", files={"document": ("deck.pdf", make_pdf())},
                           data=payload).status_code == 401
        assert client.get("/api/config", headers=headers).json()["document"]["max_document_mb"] == 1
        schema = client.get("/openapi.json").json()
        multipart = schema["paths"]["/api/videos/document"]["post"]["requestBody"]["content"]["multipart/form-data"]
        assert multipart["schema"]["required"] == ["request", "document"]
        assert post(files={"document": ("bad.txt", b"bad")}).status_code == 415
        assert post(files={"document": ("bad.pdf", b"bad")}).status_code == 422
        assert post(files={"document": ("large.pdf", b"x" * (1024 * 1024 + 1))}).status_code == 413
        assert post(files={"document": ("deck.pdf", make_pdf()), "image": ("a.png", b"a")}).status_code == 422
        assert post(files={"document": ("deck.pdf", make_pdf())},
                    data={"request": VideoRequest(prompt="A movie").model_dump_json()}).status_code == 422
        assert client.post("/api/videos/document", headers=headers, json=json.loads(payload["request"])).status_code == 422
        assert client.post("/api/videos", headers=headers, json=json.loads(payload["request"])).status_code == 422
        assert post(files={"document": ("deck.pptx", make_pptx())}).status_code == 422

        def chunked():
            yield b'--b\r\nContent-Disposition: form-data; name="document"; filename="big.pdf"\r\n\r\n'
            for _ in range(48):
                yield b"x" * 65536
            yield b"\r\n--b--\r\n"

        large = client.post("/api/videos/document", content=chunked(),
                            headers={**headers, "Content-Type": "multipart/form-data; boundary=b"})
        assert large.status_code == 413 and "content-length" not in large.request.headers
        result = post(files={"document": ("deck.pdf", make_pdf())})
        assert result.status_code == 202, result.text
        job = result.json()
        assert job["request"]["document"]["filename"] == "deck.pdf"
        assert store.path(job["id"], "input/document.pdf").read_bytes() == make_pdf()
        for name in ["state.json", "input/document.pdf", "document/source.pdf", "document/pages.json"]:
            assert client.get(f"/api/videos/{job['id']}/media/{name}", headers=headers).status_code == 404


@requires_ffmpeg
def test_document_api_completes_and_serves_script_video_and_events(setup):
    settings, store, team, narrator, factory, close = setup
    app = create_app(settings, services=(store, factory, close))
    with TestClient(app) as client:
        result = client.post("/api/videos/document", data={"request": request().model_dump_json()},
                             files={"document": ("deck.pdf", make_pdf(), "application/pdf")})
        assert result.status_code == 202, result.text
        job_id = result.json()["id"]
        for _ in range(600):
            job = client.get(f"/api/videos/{job_id}").json()
            if job["status"] in {"completed", "failed"}:
                break
            time.sleep(0.05)
        assert job["status"] == "completed", job
        script = client.get(f"/api/videos/{job_id}/storyboard").json()
        assert script["pages"][1]["narration"] == "An explanation of page 2."
        response = client.get(f"/api/videos/{job_id}/download")
        assert response.status_code == 200 and response.headers["content-type"] == "video/mp4"
        assert b"ftyp" in response.content[:32]
        events = client.get(f"/api/videos/{job_id}/events").text
        assert "event: state" in events and '"clips_done": 2' in events and "event: end" in events
        assert client.get(f"/api/videos/{job_id}/media/scenes/scene_001.mp4").status_code == 200


async def test_narration_team_is_grounded_and_has_no_search_tools(monkeypatch):
    captured = {}

    class Agent:
        def __init__(self, **kwargs):
            captured.update(kwargs)

    async def run(agent, message, output_type):
        captured["payload"] = json.loads(message)
        if output_type is ShotPrompt:
            return ShotPrompt(prompt="A cinematic illustration of the supplied facts")
        return PageNarration(title="Slide one", narration="Narrator: A concise explanation.")

    monkeypatch.setattr(documents, "Agent", Agent)
    monkeypatch.setattr(documents, "_run_structured", run)
    team = FoundryDocumentTeam(object())
    page = documents.DocumentPage(number=1, text="Facts", notes="Emphasize the key point")
    result = await team.narrate(page, "Explain clearly", "en-US", 20)
    assert result.narration == "A concise explanation."
    assert "tools" not in captured and "untrusted" in captured["instructions"]
    assert captured["payload"]["page"]["notes"] == page.notes
    shot = await team.cutaway([page], [result.narration], "Explain clearly", "Model-specific guidance")
    assert shot.prompt and captured["name"] == "document-cinematic-director"
    assert captured["payload"]["model_guide"] == "Model-specific guidance"
    assert captured["payload"]["pages"][0]["text"] == "Facts"
    assert "tools" not in captured


@pytest.mark.parametrize("seconds,expected", [
    (15, []), (30, []), (34.99, []), (35, [30]), (60, [30]), (65, [30, 60]),
    (95, [30, 60, 90]), (600, list(range(30, 600, 30))),
])
def test_cutaway_schedule(seconds, expected):
    for fps in (12, 24, 30):
        assert documents.cutaway_starts(seconds, fps) == expected


def test_cutaways_preserve_legacy_document_serialization():
    old = request().model_dump(mode="json")
    assert "cinematic_cutaways" not in old["document"]
    assert request(document={"cinematic_cutaways": False}).model_dump(mode="json") == old
    with pytest.raises(ValidationError):
        request(document={"cinematic_cutaways": "false"})


def test_animations_preserve_legacy_document_serialization():
    old = request().model_dump(mode="json")
    assert "animations" not in old["document"]
    assert request(document={"animations": False}).model_dump(mode="json") == old
    assert request(document={"animations": True}).model_dump(mode="json")["document"]["animations"] is True
    with pytest.raises(ValidationError):
        request(document={"animations": "true"})


@requires_ffmpeg
@pytest.mark.parametrize("width,height", [(160, 90), (90, 160)])
async def test_animated_still_video_keeps_exact_duration_and_never_crops(tmp_path, width, height):
    image = tmp_path / "page.png"
    Image.new("RGB", (width, height), (200, 0, 0)).save(image)
    plain, animated = tmp_path / "plain.mp4", tmp_path / "animated.mp4"
    await media.still_video(image, plain, width, height, 12, 7.5)
    await media.still_video(image, animated, width, height, 12, 7.5, animate=True)
    assert await media.video_size(animated) == (width, height)
    assert abs(await media.video_duration(animated) - 7.5) < 1 / 12 + 0.001
    assert abs(await media.video_duration(plain) - 7.5) < 1 / 12 + 0.001
    corners = {}
    for stamp in (0, 3.0, 7.0):
        frame = tmp_path / f"frame_{stamp}.png"
        await media.run("ffmpeg", "-y", "-ss", str(stamp), "-i", str(animated), "-frames:v", "1", str(frame))
        with Image.open(frame) as shot:
            corners[stamp] = (shot.getpixel((0, 0)), shot.getpixel((width // 2, height // 2)))
    # Fades start from white; mid-page the colour is full strength; the white margin shrinks as the page zooms in.
    assert corners[0][1][1] > 200
    assert corners[3.0][1][0] > 180 and corners[3.0][1][1] < 40
    assert corners[3.0][0][1] > 150
    assert corners[7.0][0][1] < corners[3.0][0][1]


@requires_ffmpeg
async def test_document_pipeline_with_animations(setup):
    settings, store, team, narrator, factory, close = setup
    upload = validate_document(make_pdf(), "deck.pdf", settings.max_document_bytes)
    manager = JobManager(settings, store, factory)
    state = await manager.create(request(document={"animations": True}), document=upload)
    await manager.wait(state.id)
    finished = await manager.get(state.id)
    assert finished.status == "completed", finished.error
    final = store.path(state.id, "final.mp4")
    seconds, has_audio = await media.probe(final)
    assert has_audio and abs(seconds - 15) < 0.1
    assert await media.video_size(final) == (160, 90)
    for i, stamp in enumerate([3, 11]):
        frame = final.parent / f"anim_{i}.png"
        await media.run("ffmpeg", "-y", "-ss", str(stamp), "-i", str(final), "-frames:v", "1", str(frame))
        with Image.open(frame) as image:
            r, g, b = image.getpixel((image.width // 2, image.height // 2))
            assert (r > 180 and b < 60) if i == 0 else (b > 180 and r < 60)
        frame.unlink()
    await manager.shutdown()


@requires_ffmpeg
@pytest.mark.parametrize("fail_first,fps,orientation", [
    (False, 24, "horizontal"), (True, 12, "horizontal"), (False, 24, "vertical"),
])
async def test_cinematic_cutaways_exact_windows_voice_and_retry(setup, fail_first, fps, orientation):
    settings, store, _, narrator, _, _ = setup
    settings = replace(settings, output_fps=fps)
    prompts, generations = [], []

    class Team(DocumentTeam):
        async def cutaway(self, pages, narration, prompt, guide):
            prompts.append([page.number for page in pages])
            return ShotPrompt(prompt="A cinematic green scene illustrating the document")

    class Comfy:
        async def generate_clip(self, **kwargs):
            generations.append(kwargs["prompt"])
            if fail_first and len(generations) == 2:
                assert kwargs["resume"] == {"server": "http://test-comfy", "prompt_id": "prompt-id"}
            await kwargs["on_submitted"]("http://test-comfy", "prompt-id")
            if fail_first and len(generations) == 1:
                raise RuntimeError("GPU unavailable")
            dest = kwargs["dest"]
            dest.parent.mkdir(parents=True, exist_ok=True)
            await media.run("ffmpeg", "-y", "-f", "lavfi", "-i", "color=lime:s=160x90:r=12:d=6",
                            "-f", "lavfi", "-i", "sine=frequency=880:duration=6",
                            *media.VIDEO_ENCODE, *media.AUDIO_ENCODE, str(dest))

    team = Team()

    def factory(manager):
        return build_document_workflow(PipelineDeps(
            settings=settings, team=None, comfy=Comfy(), narrator=narrator, store=store,
            progress=manager.progress, ops=manager.recorder,
        ), team)

    manager = JobManager(settings, store, factory)
    req = request(document={"cinematic_cutaways": True}, orientation=orientation).model_copy(
        update={"duration_minutes": 65 / 60})
    state = await manager.create(req, document=validate_document(make_pdf(), "deck.pdf", 10000))
    await manager.wait(state.id)
    if fail_first:
        failed = await manager.get(state.id)
        assert failed.status == "failed" and "GPU unavailable" in failed.error
        assert not await store.exists(state.id, "final.mp4")
        await manager.retry(state.id)
        await manager.wait(state.id)
    finished = await manager.get(state.id)
    assert finished.status == "completed", finished.error
    assert prompts == [[1, 2], [2]]
    assert len(generations) == (3 if fail_first else 2)
    final, base = store.path(state.id, "final.mp4"), store.path(state.id, "document/base.mp4")
    assert await media.video_duration(final) == 65
    assert await media.video_size(final) == ((160, 90) if orientation == "horizontal" else (90, 160))
    for stamp, cinematic in [(29.75, False), (30, True), (34.75, True), (35, False),
                             (59.75, False), (60, True), (64.75, True)]:
        frame = final.parent / "verify.png"
        await media.run("ffmpeg", "-y", "-ss", str(stamp), "-i", str(final), "-frames:v", "1", str(frame))
        with Image.open(frame) as image:
            r, g, b = image.getpixel((image.width // 2, image.height // 2))
            assert (g > 220 and r < 30 and b < 30) == cinematic
        frame.unlink()
    audio = []
    for path in (base, final):
        result = subprocess.run(["ffmpeg", "-v", "error", "-i", str(path), "-map", "0:a:0",
                                 "-c:a", "copy", "-f", "adts", "-"], capture_output=True, check=True)
        audio.append(result.stdout)
    assert audio[0] == audio[1]
    await manager.progress(state.id, JobStatus.failed, error="Interrupted after rendering")
    await manager.retry(state.id)
    await manager.wait(state.id)
    assert (await manager.get(state.id)).status == "completed"
    assert len(prompts) == 2 and len(narrator.texts) == 2
    assert len(generations) == (3 if fail_first else 2)
    await manager.shutdown()


@pytest.mark.parametrize("outcome", ["success", "failure", "missing", "timeout"])
async def test_powerpoint_conversion_process_and_failures(tmp_path, monkeypatch, outcome):
    source = tmp_path / "document.pptx"
    source.write_bytes(make_pptx())
    settings = Settings(libreoffice_path="office-test")
    monkeypatch.setattr(documents, "office_available", lambda settings: True)
    calls = []

    class Process:
        returncode = None
        killed = False

        async def communicate(self):
            if outcome == "timeout":
                await asyncio.sleep(1)
            self.returncode = 1 if outcome == "failure" else 0
            return b"converted", b"Conversion error" if self.returncode else b""

        def kill(self):
            self.killed = True
            self.returncode = -9

        async def wait(self):
            return self.returncode

    proc = Process()

    async def start(*args, **kwargs):
        calls.append(args)
        if outcome == "success":
            dest = tmp_path / "converted" / "document.pdf"
            dest.write_bytes(make_pdf())
        return proc

    monkeypatch.setattr(documents.asyncio, "create_subprocess_exec", start)
    if outcome == "timeout":
        wait_for = asyncio.wait_for

        async def short_wait(awaitable, timeout):
            return await wait_for(awaitable, 0.01)

        monkeypatch.setattr(documents.asyncio, "wait_for", short_wait)
    if outcome == "success":
        pdf = await documents.convert_powerpoint(source, settings)
        assert [p.text for p in pdf_pages(pdf)] == ["First page", "Second page"]
    else:
        with pytest.raises(DocumentError, match="timed out" if outcome == "timeout" else "conversion failed"):
            await documents.convert_powerpoint(source, settings)
    assert calls[0][0] == "office-test" and "--headless" in calls[0]
    assert "ExportHiddenSlides" in calls[0][calls[0].index("--convert-to") + 1]
    assert any(arg.startswith("-env:UserInstallation=file:") for arg in calls[0])
    assert proc.killed == (outcome == "timeout")
